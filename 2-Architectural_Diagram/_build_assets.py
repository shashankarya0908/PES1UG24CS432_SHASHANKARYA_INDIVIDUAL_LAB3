from __future__ import annotations

import html
import textwrap
import xml.etree.ElementTree as ET
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from reportlab.graphics import renderPDF, renderPM
from reportlab.graphics.shapes import Circle, Drawing, Line, Path as ShapePath, Rect, String
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

ROOT = Path(__file__).resolve().parents[1]
OUT = Path(__file__).resolve().parent
W, H = 1800, 1240

components = [
    dict(id="portal", name="Showcase Portal", role="Presentation boundary", x=700, y=1030, w=400, h=120, layer="presentation", provides=[], requires=["IShowcaseCatalog", "IProjectSubmission", "IRubricEvaluation", "IPublicVoting", "ILeaderboard"]),
    dict(id="profile", name="Showcase Profile", role="Team profile and published project catalog", x=80, y=820, w=360, h=140, layer="business", provides=["IShowcaseCatalog"], requires=["IProjectData"]),
    dict(id="submission", name="Project Submission", role="Upload workflow and artifact references", x=520, y=820, w=360, h=140, layer="business", provides=["IProjectSubmission"], requires=["IShowcaseCatalog", "IProjectData", "IArtifactStore"]),
    dict(id="evaluation", name="Evaluation & Rubric", role="Judge scores, comments, and results", x=1040, y=820, w=360, h=140, layer="business", provides=["IRubricEvaluation"], requires=["IShowcaseCatalog", "IProjectData"]),
    dict(id="identity", name="Identity Verification", role="Verifies voter and returns verified reference", x=80, y=590, w=360, h=140, layer="business", provides=["IIdentityVerification"], requires=[]),
    dict(id="voting", name="Public Voting", role="Authorizes vote and enforces one per project", x=520, y=590, w=360, h=140, layer="business", provides=["IPublicVoting", "IAcceptedVoteEvents"], requires=["IIdentityVerification", "IShowcaseCatalog", "IVoteLedger"]),
    dict(id="leaderboard", name="Leaderboard & Realtime", role="Aggregates counts and publishes ranking updates", x=1040, y=590, w=360, h=140, layer="business", provides=["ILeaderboard"], requires=["IShowcaseCatalog", "IVoteLedger", "IAcceptedVoteEvents"]),
    dict(id="project_data", name="Project Data Repository", role="Profiles, projects, artifact refs, evaluations", x=60, y=230, w=440, h=150, layer="data", provides=["IProjectData"], requires=[]),
    dict(id="vote_ledger", name="Vote Ledger", role="Atomic unique voter/project vote records", x=680, y=230, w=440, h=150, layer="data", provides=["IVoteLedger"], requires=[]),
    dict(id="artifact_store", name="Multimedia Artifact Store", role="Uploaded media objects and artifact references", x=1260, y=230, w=480, h=150, layer="data", provides=["IArtifactStore"], requires=[]),
]

# Each connector starts at its requiring socket and ends at its provider lollipop.
edges = [
    dict(name="IShowcaseCatalog", source="portal", target="profile", pts=[(700, 1090), (30, 1090), (30, 990), (260, 990), (260, 960)], label=(110, 1002)),
    dict(name="IProjectSubmission", source="portal", target="submission", pts=[(790, 1030), (790, 995), (700, 995), (700, 960)], label=(795, 1002)),
    dict(name="IRubricEvaluation", source="portal", target="evaluation", pts=[(1010, 1030), (1010, 995), (1220, 995), (1220, 960)], label=(1020, 1002)),
    dict(name="IPublicVoting", source="portal", target="voting", pts=[(700, 1120), (25, 1120), (25, 650), (520, 650)], label=(40, 680)),
    dict(name="ILeaderboard", source="portal", target="leaderboard", pts=[(1100, 1120), (1770, 1120), (1770, 650), (1400, 650)], label=(1580, 680)),
    dict(name="IIdentityVerification", source="voting", target="identity", pts=[(520, 650), (440, 650)], label=(414, 676)),
    dict(name="IShowcaseCatalog", source="submission", target="profile", pts=[(520, 875), (440, 875)], label=(421, 897)),
    dict(name="IShowcaseCatalog", source="evaluation", target="profile", pts=[(1040, 900), (1000, 900), (1000, 980), (220, 980), (220, 960)], label=None),
    dict(name="IShowcaseCatalog", source="voting", target="profile", pts=[(520, 700), (480, 700), (480, 975), (300, 975), (300, 960)], label=None),
    dict(name="IShowcaseCatalog", source="leaderboard", target="profile", pts=[(1040, 700), (1015, 700), (1015, 1005), (380, 1005), (380, 960)], label=None),
    dict(name="IProjectData", source="profile", target="project_data", pts=[(260, 820), (260, 380)], label=(270, 470)),
    dict(name="IProjectData", source="submission", target="project_data", pts=[(520, 875), (480, 875), (480, 430), (500, 430), (500, 305)], label=None),
    dict(name="IProjectData", source="evaluation", target="project_data", pts=[(1400, 875), (1740, 875), (1740, 440), (480, 440), (480, 305)], label=None),
    dict(name="IArtifactStore", source="submission", target="artifact_store", pts=[(880, 875), (900, 875), (900, 420), (1500, 420), (1500, 380)], label=(1110, 425)),
    dict(name="IVoteLedger", source="voting", target="vote_ledger", pts=[(700, 590), (700, 480), (900, 480), (900, 380)], label=(710, 490)),
    dict(name="IVoteLedger", source="leaderboard", target="vote_ledger", pts=[(1200, 590), (1200, 470), (900, 470), (900, 380)], label=None),
    dict(name="IAcceptedVoteEvents", source="leaderboard", target="voting", pts=[(1040, 650), (880, 650)], label=(883, 674)),
]

INK = colors.HexColor("#23323A")
MUTED = colors.HexColor("#53666D")
STROKE = colors.HexColor("#53666D")
BLUE = colors.HexColor("#39738A")
GREEN = colors.HexColor("#477B63")
GOLD = colors.HexColor("#A17A29")
RED = colors.HexColor("#A65F4A")
WHITE = colors.white


def add_text(d, x, y, text, size=10, color=INK, bold=False):
    d.add(String(x, y, text, fontName="Helvetica-Bold" if bold else "Helvetica", fontSize=size, fillColor=color))


def add_component(d, c):
    x, y, w, h = c["x"], c["y"], c["w"], c["h"]
    outline = BLUE if c["layer"] == "presentation" else GREEN
    d.add(Rect(x, y, w, h, rx=4, ry=4, fillColor=WHITE, strokeColor=outline, strokeWidth=2))
    d.add(Rect(x+w-31, y+h-22, 14, 7, fillColor=WHITE, strokeColor=outline, strokeWidth=1.4))
    d.add(Rect(x+w-31, y+h-33, 14, 7, fillColor=WHITE, strokeColor=outline, strokeWidth=1.4))
    add_text(d, x+13, y+h-27, c["name"], 14, INK, True)
    add_text(d, x+13, y+h-45, c["role"], 9, MUTED)
    add_text(d, x+13, y+h-68, "Provides", 8.5, outline, True)
    provided = ", ".join(c["provides"]) or "-"
    add_text(d, x+70, y+h-68, provided, 8.2, INK)
    add_text(d, x+13, y+h-86, "Requires", 8.5, RED, True)
    required = ", ".join(c["requires"]) or "-"
    chunks = textwrap.wrap(required, width=max(28, int((w-92)/4.5))) or ["-"]
    for i, chunk in enumerate(chunks[:2]):
        add_text(d, x+70, y+h-86-i*12, chunk, 7.4, INK)


def add_socket(d, x, y, dx, dy):
    r = 7
    p = ShapePath(strokeColor=INK, strokeWidth=2, fillColor=None)
    if dx < 0:
        p.moveTo(x, y-r); p.curveTo(x-r, y-r, x-r, y+r, x, y+r)
    elif dx > 0:
        p.moveTo(x, y-r); p.curveTo(x+r, y-r, x+r, y+r, x, y+r)
    elif dy > 0:
        p.moveTo(x-r, y); p.curveTo(x-r, y+r, x+r, y+r, x+r, y)
    else:
        p.moveTo(x-r, y); p.curveTo(x-r, y-r, x+r, y-r, x+r, y)
    d.add(p)


def add_edge(d, e):
    pts = e["pts"]
    for (x1, y1), (x2, y2) in zip(pts, pts[1:]):
        d.add(Line(x1, y1, x2, y2, strokeColor=STROKE, strokeWidth=1.6))
    sx, sy = pts[0]
    ex, ey = pts[-1]
    add_socket(d, sx, sy, pts[1][0]-sx, pts[1][1]-sy)
    d.add(Circle(ex, ey, 6.5, fillColor=WHITE, strokeColor=INK, strokeWidth=1.8))
    lx, ly = e["label"]
    width = max(85, len(e["name"])*5.5)
    d.add(Rect(lx-3, ly-3, width, 15, fillColor=WHITE, strokeColor=None))
    add_text(d, lx, ly, e["name"], 8, INK, True)


def write_diagram():
    d = Drawing(W, H + 40)
    d.add(Rect(0, 0, W, H, fillColor=WHITE, strokeColor=None))
    bands = [(40, 1020, 1720, 185, colors.HexColor("#EEF3F5"), BLUE, "PRESENTATION LAYER"), (40, 545, 1720, 435, colors.HexColor("#F4F6F4"), GREEN, "APPLICATION / BUSINESS LAYER"), (40, 195, 1720, 350, colors.HexColor("#F8F4E9"), GOLD, "DATA LAYER")]
    for x,y,w,h,fill,outline,label in bands:
        d.add(Rect(x,y,w,h,fillColor=fill,strokeColor=colors.HexColor("#CFD8D9"),strokeWidth=1))
        add_text(d,x+14,y+h-22,label,11,outline,True)
    add_text(d, 42, 1250, "Student Project Portfolio & Showcase App", 25, INK, True)
    add_text(d, 42, 1220, "UML Component Architecture  |  Shashank Arya  |  PES1UG24CS432  |  Problem Statement #04", 14, MUTED)
    for edge in edges:
        pts=edge["pts"]
        for (x1,y1),(x2,y2) in zip(pts,pts[1:]):
            d.add(Line(x1,y1,x2,y2,strokeColor=STROKE,strokeWidth=1.6))
    for comp in components:
        add_component(d,comp)
    # Port marks and labels are painted above the boxes so the UML symbols stay visible.
    for edge in edges:
        pts=edge["pts"]
        sx,sy=pts[0]; ex,ey=pts[-1]
        add_socket(d,sx,sy,pts[1][0]-sx,pts[1][1]-sy)
        d.add(Circle(ex,ey,6.5,fillColor=WHITE,strokeColor=INK,strokeWidth=1.8))
        if edge["label"] is None:
            continue
        lx,ly=edge["label"]
        width=max(85,len(edge["name"])*5.5)
        d.add(Rect(lx-3,ly-3,width,15,fillColor=WHITE,strokeColor=None))
        add_text(d,lx,ly,edge["name"],8,INK,True)
    add_text(d,60,150,"Assembly connectors pair a required socket at the consumer with a provided lollipop at the provider.",10,INK)
    add_text(d,60,128,"Vote flow: identity verification -> atomic unique (verifiedVoterRef, projectId) commit -> VoteAccepted event -> refreshed live ranking.",10,INK)
    add_text(d,60,106,"Ledger uniqueness supports one vote per project; identity verification supports NFR-002 but does not guarantee elimination of every sybil attack.",9,MUTED)
    renderPDF.drawToFile(d,str(OUT/"Architecture_Diagram.pdf"))
    renderPM.drawToFile(d,str(OUT/"Architecture_Diagram.png"),fmt="PNG",dpi=150)


def xml_cell(parent, cid, value, style, x=None, y=None, w=None, h=None, edge=False, source=None, target=None, points=None):
    cell=ET.SubElement(parent,"mxCell",{"id":cid,"value":value,"style":style,"parent":"1"})
    cell.set("edge" if edge else "vertex","1")
    if source: cell.set("source",source)
    if target: cell.set("target",target)
    geo=ET.SubElement(cell,"mxGeometry",{"relative":"1" if edge else "0","as":"geometry"})
    if not edge:
        geo.attrib.update({"x":str(x),"y":str(y),"width":str(w),"height":str(h)})
    elif points:
        arr=ET.SubElement(geo,"Array",{"as":"points"})
        for px,py in points[1:-1]: ET.SubElement(arr,"mxPoint",{"x":str(px),"y":str(H-py)})
    return cell


def write_drawio():
    mx=ET.Element("mxfile",{"host":"app.diagrams.net","modified":"2026-10-02T00:00:00.000Z","agent":"GitHub Copilot","version":"24.7.17","type":"device"})
    dg=ET.SubElement(mx,"diagram",{"id":"portfolio-component-architecture","name":"Component Architecture"})
    model=ET.SubElement(dg,"mxGraphModel",{"dx":str(W),"dy":str(H),"grid":"1","gridSize":"10","guides":"1","tooltips":"1","connect":"1","arrows":"1","fold":"1","page":"1","pageScale":"1","pageWidth":str(W),"pageHeight":str(H),"math":"0","shadow":"0"})
    root=ET.SubElement(model,"root")
    ET.SubElement(root,"mxCell",{"id":"0"}); ET.SubElement(root,"mxCell",{"id":"1","parent":"0"})
    xml_cell(root,"title","Student Project Portfolio &amp; Showcase App | UML Component Architecture","text;html=1;align=left;verticalAlign=middle;fontSize=25;fontStyle=1;fontColor=#23323A;strokeColor=none;fillColor=none;",40,15,1650,45)
    xml_cell(root,"meta","Shashank Arya | PES1UG24CS432 | Problem Statement #04","text;html=1;align=left;verticalAlign=middle;fontSize=14;fontColor=#53666D;strokeColor=none;fillColor=none;",40,60,1650,30)
    for band_id,label,x,y,w,h,fill in [("presentation","PRESENTATION LAYER",40,35,1720,185,"#EEF3F5"),("business","APPLICATION / BUSINESS LAYER",40,260,1720,435,"#F4F6F4"),("data","DATA LAYER",40,695,1720,350,"#F8F4E9")]:
        xml_cell(root,"band_"+band_id,label,f"swimlane;horizontal=1;startSize=30;rounded=0;html=1;whiteSpace=wrap;align=left;verticalAlign=top;spacingLeft=12;fontSize=14;fontStyle=1;fontColor=#53666D;fillColor={fill};swimlaneFillColor={fill};strokeColor=#CFD8D9;",x,y,w,h)
    ids={}
    for c in components:
        provided=", ".join(c["provides"]) or "-"; required=", ".join(c["requires"]) or "-"
        val=f"<b>{html.escape(c['name'])}</b><br/>{html.escape(c['role'])}<br/><b>Provides:</b> {html.escape(provided)}<br/><b>Requires:</b> {html.escape(required)}"
        style="shape=component;jettyWidth=26;jettyHeight=13;whiteSpace=wrap;html=1;align=left;verticalAlign=middle;spacing=10;fontSize=13;fontColor=#23323A;strokeColor=#477B63;strokeWidth=1.5;fillColor=#FFFFFF;"
        if c["id"]=="portal": style=style.replace("#477B63","#39738A")
        cell=xml_cell(root,c["id"],val,style,c["x"],H-c["y"]-c["h"],c["w"],c["h"])
        ids[c["id"]]=cell
    for i,e in enumerate(edges,1):
        style="edgeStyle=orthogonalEdgeStyle;rounded=0;orthogonalLoop=1;jettySize=auto;html=1;strokeColor=#53666D;strokeWidth=1.5;fontSize=10;fontColor=#23323A;labelBackgroundColor=#FFFFFF;startArrow=halfCircle;startFill=0;startSize=12;endArrow=oval;endFill=0;endSize=12;"
        xml_cell(root,f"edge_{i}",e["name"],style,edge=True,source=e["source"],target=e["target"],points=e["pts"])
    xml_cell(root,"legend","Socket = required interface | Lollipop = provided interface | Connector direction shown from requiring component to providing component. Vote ledger commit precedes VoteAccepted event and live ranking refresh.","text;html=1;align=left;verticalAlign=middle;whiteSpace=wrap;fontSize=12;fontColor=#23323A;strokeColor=none;fillColor=none;",60,H-175,1680,55)
    ET.indent(mx,space="  ")
    (OUT/"Architecture_Diagram.drawio").write_bytes(ET.tostring(mx,encoding="utf-8",xml_declaration=True))


JUSTIFICATION = [
    "The Student Project Portfolio & Showcase App combines team-owned project evidence, judge-only rubric assessment, and public voting. A layered component architecture gives these workflows a clear presentation boundary while keeping profile/submission, evaluation, voting, ranking, and persistence responsibilities separate. Project records, vote records, and multimedia objects are stored behind distinct interfaces to reflect their different integrity and handling needs.",
    "Voting is the highest-integrity boundary in the supplied requirements. Public Voting requires a verified voter reference and asks the Vote Ledger to atomically record the voter/project pair; a storage-level uniqueness rule supports FR-003 under repeated or concurrent requests. Identity Verification validates the voter rather than trusting portal-supplied credentials. This supports NFR-002 without claiming that identity checks alone eliminate every sybil attempt.",
    "For NFR-001, the leaderboard consumes accepted-vote events only after the ledger commit, refreshes totals from authoritative vote data, and publishes ranking updates to the portal. This separates vote acceptance from ranking refresh without promising an unspecified latency or throughput target. Interfaces remain technology-neutral because no provider, protocol, or deployment topology was supplied.",
    "The component boundaries directly cover FR-001 through FR-005 and the four stated project areas. They also make the identity and duplicate-vote controls visible for comparison with the Jira Bug Tracker. That comparison is pending: the official Lab 3 handout and Scrum, Kanban, and Jira source documents were not in the workspace, so no defect IDs or undocumented constraints are asserted. This is a logical architecture, not evidence of an implemented or tested system.",
]


def write_justification():
    doc=Document(); sec=doc.sections[0]
    sec.top_margin=Inches(.48); sec.bottom_margin=Inches(.48); sec.left_margin=Inches(.62); sec.right_margin=Inches(.62)
    normal=doc.styles["Normal"]; normal.font.name="Aptos"; normal.font.size=Pt(9); normal.font.color.rgb=RGBColor(35,50,58); normal.paragraph_format.space_after=Pt(5)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run("ARCHITECTURE JUSTIFICATION"); r.bold=True; r.font.name="Aptos Display"; r.font.size=Pt(16); r.font.color.rgb=RGBColor(35,50,58)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(8)
    r=p.add_run("Shashank Arya  |  PES1UG24CS432  |  Problem Statement #04 – Student Project Portfolio & Showcase App"); r.font.size=Pt(8); r.font.color.rgb=RGBColor(71,123,99)
    for text in JUSTIFICATION: doc.add_paragraph(text)
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(0)
    r=p.add_run("Source limitation: official handout and Scrum/Kanban/Jira documents were unavailable; reconcile before final upload."); r.italic=True; r.font.size=Pt(8); r.font.color.rgb=RGBColor(83,102,109)
    doc.save(OUT/"Architecture_Justification_Document.docx")
    styles=getSampleStyleSheet()
    story=[Paragraph("ARCHITECTURE JUSTIFICATION",ParagraphStyle("T",parent=styles["Title"],fontName="Helvetica-Bold",fontSize=15,leading=18,alignment=TA_CENTER,textColor=INK,spaceAfter=4)),Paragraph("Shashank Arya  |  PES1UG24CS432  |  Problem Statement #04 – Student Project Portfolio &amp; Showcase App",ParagraphStyle("M",parent=styles["Normal"],fontSize=8,leading=10,alignment=TA_CENTER,textColor=GREEN,spaceAfter=9))]
    for text in JUSTIFICATION: story.append(Paragraph(html.escape(text),ParagraphStyle("B",parent=styles["BodyText"],fontName="Helvetica",fontSize=8.7,leading=11.5,textColor=INK,spaceAfter=6)))
    story.append(Paragraph("Source limitation: official handout and Scrum/Kanban/Jira documents were unavailable; reconcile before final upload.",ParagraphStyle("N",parent=styles["BodyText"],fontName="Helvetica-Oblique",fontSize=7.7,leading=9,textColor=MUTED)))
    SimpleDocTemplate(str(OUT/"Architecture_Justification_Document.pdf"),pagesize=letter,rightMargin=.62*inch,leftMargin=.62*inch,topMargin=.48*inch,bottomMargin=.48*inch,title="Architecture Justification",author="Shashank Arya").build(story)


def md_story(text):
    styles=getSampleStyleSheet(); body=ParagraphStyle("BodyM",parent=styles["BodyText"],fontSize=8.5,leading=10.6,spaceAfter=3,textColor=INK)
    hs={1:ParagraphStyle("H1M",parent=styles["Heading1"],fontSize=17,leading=20,spaceBefore=12,spaceAfter=6,textColor=INK),2:ParagraphStyle("H2M",parent=styles["Heading2"],fontSize=12.5,leading=15,spaceBefore=9,spaceAfter=4,textColor=GREEN),3:ParagraphStyle("H3M",parent=styles["Heading3"],fontSize=10.5,leading=13,spaceBefore=7,spaceAfter=3,textColor=INK)}
    story=[]; lines=text.splitlines(); i=0
    while i<len(lines):
        line=lines[i].strip()
        if not line: i+=1; continue
        if line.startswith("|"):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith("|"):
                row=[v.strip() for v in lines[i].strip().strip("|").split("|")]
                if not all(set(v.replace(":", ""))<={"-"} for v in row): rows.append(row)
                i+=1
            if rows:
                n=max(len(r) for r in rows); data=[[Paragraph(html.escape(v),body) for v in r+[""]*(n-len(r))] for r in rows]
                t=Table(data,colWidths=[7.25*inch/n]*n,repeatRows=1)
                t.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#E2EEE8")),("GRID",(0,0),(-1,-1),.35,colors.HexColor("#B7C4C2")),("VALIGN",(0,0),(-1,-1),"TOP"),("LEFTPADDING",(0,0),(-1,-1),3),("RIGHTPADDING",(0,0),(-1,-1),3),("TOPPADDING",(0,0),(-1,-1),2),("BOTTOMPADDING",(0,0),(-1,-1),2)]))
                story.extend([t,Spacer(1,4)])
            continue
        match=__import__("re").match(r"^(#{1,3})\s+(.*)$",line)
        if match: story.append(Paragraph(html.escape(match.group(2)),hs[len(match.group(1))])); i+=1; continue
        if line.startswith("- "): story.append(Paragraph("&#8226; "+html.escape(line[2:]),body)); i+=1; continue
        story.append(Paragraph(html.escape(line).replace("`",""),body)); i+=1
    return story


def write_spec_pdf():
    spec=(OUT/"Architecture_Specification.md").read_text(encoding="utf-8")
    SimpleDocTemplate(str(OUT/"Architecture_Specification.pdf"),pagesize=letter,rightMargin=.62*inch,leftMargin=.62*inch,topMargin=.5*inch,bottomMargin=.5*inch,title="Architecture Specification",author="Shashank Arya").build(md_story(spec))


if __name__=="__main__":
    write_diagram(); write_drawio(); write_justification(); write_spec_pdf()
    print("Created draw.io XML, PNG/PDF diagram exports, DOCX/PDF justification, and specification PDF.")